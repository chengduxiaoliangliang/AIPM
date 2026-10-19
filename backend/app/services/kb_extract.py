"""Text extraction for knowledge-base sources.

Turns an uploaded file into plain text so the LLM can read it. Every extractor
degrades gracefully: on failure we return a note instead of raising, and the
source is still stored and downloadable.
"""
from __future__ import annotations

import html
import io
import re
from pathlib import Path
from typing import Tuple

# Anything we can turn into readable text.
TEXT_LIKE = {
    ".md", ".markdown", ".txt", ".text", ".log", ".csv", ".tsv",
    ".json", ".yaml", ".yml", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".vue", ".java", ".go", ".sql", ".sh",
    ".ini", ".conf", ".toml", ".env", ".properties",
}
SUPPORTED = TEXT_LIKE | {".pdf", ".docx", ".xlsx", ".xls"}

MAX_CHARS = 300_000
_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "gbk", "big5", "latin-1")


def is_supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED


def _decode(raw: bytes) -> str:
    for enc in _ENCODINGS:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def _strip_html(text: str) -> str:
    text = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    text = re.sub(r"(?i)</(p|div|li|tr|h[1-6])>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _extract_pdf(raw: bytes) -> Tuple[str, str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "", "未安装 pypdf，无法解析 PDF"
    reader = PdfReader(io.BytesIO(raw))
    parts = []
    for i, page in enumerate(reader.pages):
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    text = "\n\n".join(p for p in parts if p.strip())
    note = f"PDF {len(reader.pages)} 页"
    if not text.strip():
        note += "；未能提取到文字（可能是扫描件，需 OCR）"
    return text, note


def _extract_docx(raw: bytes) -> Tuple[str, str]:
    try:
        import docx
    except ImportError:
        return "", "未安装 python-docx，无法解析 Word"
    doc = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts), f"Word 文档，{len(doc.paragraphs)} 段落"


def _extract_xlsx(raw: bytes) -> Tuple[str, str]:
    try:
        import openpyxl
    except ImportError:
        return "", "未安装 openpyxl，无法解析 Excel"
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## 工作表: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts), f"Excel，{len(wb.sheetnames)} 个工作表"


def extract_text(path: str | Path, original_name: str = "") -> Tuple[str, str]:
    """Return (text, note). Never raises — returns an explanatory note on failure."""
    p = Path(path)
    ext = Path(original_name or p.name).suffix.lower()

    try:
        raw = p.read_bytes()
    except Exception as exc:
        return "", f"读取文件失败: {exc}"

    try:
        if ext == ".pdf":
            text, note = _extract_pdf(raw)
        elif ext == ".docx":
            text, note = _extract_docx(raw)
        elif ext in (".xlsx", ".xls"):
            text, note = _extract_xlsx(raw)
        elif ext in (".html", ".htm"):
            text, note = _strip_html(_decode(raw)), "HTML 已去除标签"
        elif ext in TEXT_LIKE:
            text, note = _decode(raw), "纯文本"
        else:
            text, note = _decode(raw), f"未识别的格式 {ext}，按纯文本尝试读取"
    except Exception as exc:
        return "", f"解析失败({ext}): {exc}"

    text = text.replace("\x00", " ").strip()
    if len(text) > MAX_CHARS:
        note = f"{note}；内容过长，已截断至 {MAX_CHARS} 字符"
        text = text[:MAX_CHARS]

    if not text:
        return "", note or "未提取到文本内容"
    return text, note
